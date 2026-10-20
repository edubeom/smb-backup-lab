#!/usr/bin/env python3
"""Build the offline lecture, narrated slide deck and broadcast handout bundle."""
from datetime import datetime
import hashlib
import html
import json
from pathlib import Path
import shutil
import sys
import zipfile
from zoneinfo import ZoneInfo

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'media/broadcast/2026-10-20'
REPO='https://github.com/sungbeom78/smb-backup-lab'
UPSTREAM='https://github.com/sungbeom78/dev-process-light'
PIN='4bca879961e5a5d1fda52acaf2ecf85fce6f428f'
DECK=[]

def cards(items):
    return '<div class="cards">'+''.join(f'<div class="card"><b>{a}</b><span>{b}</span></div>' for a,b in items)+'</div>'

def flow(items):
    return '<ol class="flow">'+''.join(f'<li><b>{a}</b><span>{b}</span></li>' for a,b in items)+'</ol>'

def screen(title,body):
    return f'<div class="screen"><div class="top"><span class="t">{title}</span></div><div class="in">{body}</div></div>'

def add(title,chapter,body,notes,seconds=45,kind='concept',image=None):
    DECK.append(dict(title=title,chapter=chapter,body=body,notes=notes,seconds=seconds,kind=kind,image=image))


add('AI와 만드는 서버 백업<br>기록이 복구의 근거가 되다','INTRO',
    '<div class="hero"><p class="eyebrow">dev-process-light × smb-backup-lab</p><p>요구 정리부터 오류 재현·백업·복구 확인까지</p><div class="tag">실제 개발 사례 · 약 20분 · 16:9 방송 자료</div></div>',
    '안녕하세요. 오늘은 AI와 함께 SMB 공유에 서버를 백업하는 도구를 만든 과정을 소개합니다. 중심 주제는 dev-process-light입니다. 코드를 빨리 만드는 장면보다, 무엇을 원했고 무엇을 확인했으며 다음 사람이 어디서 이어갈 수 있는지 기록하는 과정을 보여 드리겠습니다. 실제 서비스를 유지하며 진행한 사례이고, 확인한 결과와 아직 남은 시험을 나눠 설명하겠습니다.',40,'cover')
add('오늘 가져갈 세 가지','INTRO',cards([
    ('작업을 작게','목표·완료 기준·영향 범위를 한 장에'),
    ('성공을 구체적으로','파일 생성·복원·서비스 검증을 구분'),
    ('다음 작업에 연결','기록·Git 저장·단계 스냅샷을 남기기')]),
    '이 강의에서 가져갈 것은 세 가지입니다. 첫째는 작업을 작은 단위로 정리하는 방법입니다. 둘째는 성공했다는 말 대신 실제 검증 범위를 적는 방법입니다. 셋째는 기록과 Git, 중간 화면을 다음 작업의 출발점으로 사용하는 방법입니다. 백업 경험이 적어도 흐름을 따라올 수 있도록 실제 사례를 단순한 그림으로 설명하겠습니다.',30,'overview')
add('시작은 “어디에 무엇이 있지?”','WHY',cards([
    ('프로젝트','코드 · Git · 환경 설정 · 업로드'),('서버','SSH · 터널 · 자동 실행 설정'),('데이터','PostgreSQL · 메일 · SQLite')]),
    '처음부터 압축 명령을 실행하지 않았습니다. 서버에는 프로젝트 외에도 데이터베이스와 메일, SSH와 터널 설정이 있었습니다. 실제 위치와 기존 백업을 먼저 조사해야 복구할 때 빠뜨리지 않습니다. 이 단계의 원본 조사에는 서버 식별 정보가 들어가므로 방송에서는 지금처럼 범주만 보여 줍니다. 조사와 공개 자료를 구분하는 것도 기록의 한 부분입니다.',45)
add('dev-process-light는 작업 방법을 남깁니다','FRAMEWORK',cards([
    ('규칙','AGENTS.md · 무엇을 먼저 확인할지'),('스킬','light-* · 단계별 진행 방법'),('기록','dev-notes · 지금 상황과 근거')]),
    'dev-process-light는 AI와 작업할 때 따를 규칙과 스킬, 기록 형식을 프로젝트 안에 둡니다. 새 대화에서는 현재 상황을 읽고, 만들기 전에는 목표와 완료 기준을 적고, 만들고 나서는 직접 확인한 범위를 기록합니다. 이번 백업 도구는 이 방법을 실제로 적용한 사례입니다. 사용한 버전을 고정해 두었으며 원본 저장소의 최신 내용과 차이가 있을 수 있다는 점도 자료에 적었습니다.',50)
add('일곱 스킬을 필요할 때 사용합니다','FRAMEWORK',
    '<table class="grid"><tr><th>시작과 합의</th><th>확인과 공유</th><th>이어가기와 복구</th></tr><tr><td>light-resume<br>이전 기록 읽기</td><td>light-check<br>직접 확인하기</td><td>light-portal<br>기록 한눈에 보기</td></tr><tr><td>light-start<br>작업 카드 만들기</td><td>light-save<br>기록하고 Git 저장</td><td>light-undo<br>확인받고 되돌리기</td></tr><tr><td colspan="3">light-manual · 사용설명서와 강의 대본을 함께 고치기</td></tr></table>',
    '설치한 light 스킬은 일곱 가지입니다. 모두를 매번 실행하는 것은 아닙니다. 작업 시작과 합의, 확인과 저장을 중심으로 필요한 설명서와 기록 보기를 이어갑니다. 이 사례에서는 기록 포털을 따로 만들지 않았고 되돌리기도 실행하지 않았습니다. data-light는 별도의 짝 스킬로 설치했지만 새 업무 데이터 모델을 만드는 일이 아니어서 적용하지 않았습니다. 백업 메타데이터와 업무 데이터 설계를 구분했습니다.',55)
add('시작 → 합의 → 확인 → 저장','FRAMEWORK',flow([
    ('시작','NOW와 기능 지도를 읽기'),('합의','목표와 완료 기준을 정하기'),('확인','실행 결과와 한계를 적기'),('저장','기록과 Git을 함께 남기기')]),
    '전체 흐름은 네 단계로 기억하면 됩니다. 시작할 때 읽고, 만들기 전에 합의하고, 확인한 뒤에 저장합니다. 기존에 사용자가 승인한 범위는 같은 질문을 반복하지 않고 이어갑니다. 반면 서비스 중단이나 기존 백업 삭제처럼 별도 확인이 필요한 경계는 계속 지킵니다. 작은 수정부터 이 흐름을 반복하면 다음 세션에서 이전 결정을 추측할 필요가 줄어듭니다.',45)
add('요구를 작업 카드로 바꿉니다','PLAN',screen('작업 카드 · 공개용 요약',
    '<div class="checks"><p>✓ 매일 한 번 · 성공한 최근 5일 보관</p><p>✓ 프로젝트별 파일 · SSH와 터널 설정 포함</p><p>✓ 메일 서비스 계속 실행 · 비밀은 Git 밖</p><p>✓ 별도 환경에서 복구를 실제로 확인</p></div>'),
    '요구를 그대로 긴 대화에만 두면 무엇이 완료인지 흐려집니다. 그래서 목표, 눈으로 확인할 완료 기준, 영향받을 기능과 설명서를 작업 카드로 정리했습니다. 사용자의 매일 한 번, 최근 성공 다섯 개, 비밀은 Git 밖이라는 요구를 실패 조건과 보관 조건으로 옮겼습니다. 이미 승인한 진행 범위를 기록한 덕분에 같은 허락을 계속 묻는 것도 피할 수 있습니다.',55,'step')
add('기록은 세 층으로 읽습니다','PLAN',flow([
    ('LOG','전체 흐름 한 줄'),('목록','기능·날짜별 작업 번호'),('자세히','수정 파일·확인·오류·해결')])+
    '<div class="tag">NOW = 지금 상황 · FEATURES = 고칠 곳 지도 · DECISIONS = 선택의 이유</div>',
    '기록이 길어져도 전부 다시 읽지 않도록 세 층으로 나눴습니다. 먼저 전체 흐름에서 필요한 날짜를 찾고, 목록에서 작업 번호를 보고, 필요한 상세만 엽니다. 지금 상태는 NOW, 기능과 파일 연결은 FEATURES, 선택의 이유는 DECISIONS에 남깁니다. 이 자료를 만드는 시점에도 이전 작업을 이 순서로 읽고 이어갔습니다. 기록을 많이 남기는 것과 찾기 쉽게 남기는 것은 함께 가야 합니다.',45)
add('18개 단계가 한 흐름으로 남았습니다','PROCESS',
    '<div class="timeline"><div><b>01–06</b><span>규칙 · 기반 도구 · 시험</span></div><div><b>07–10</b><span>공유 연결 · SMB 오류 해결</span></div><div><b>11–12</b><span>프로젝트 분리 · 무중단 메일</span></div><div><b>13–15</b><span>첫 백업 · DB 복원 · 외부 키</span></div><div><b>16–18</b><span>예약 · 설정 호환 · 입력 개선</span></div></div>',
    '처음 여섯 단계에서는 규칙을 설치하고 기반 도구를 만들었습니다. 이후 공유기 연결 오류를 해결하고, 프로젝트별 파일과 중단 없는 메일 수집으로 요구를 구체화했습니다. 첫 실제 백업 뒤에는 데이터베이스 복원과 외부 키 확인을 진행했습니다. 마지막에는 자동 실행과 설정 호환, 승인 입력의 사용성까지 다뤘습니다. 각 화면은 그 단계에서 만든 공개용 진행 카드이며 실제 서버 화면 녹화와 구분합니다.',45)
add('SMB 공유에 암호화 파일을 보관합니다','BACKUP',flow([
    ('Linux 서버','프로젝트 · 설정 · DB · 메일'),('전용 수집','온라인 export와 보호된 임시 공간'),('암호화','프로젝트·구성요소별 GnuPG'),('SMB 공유','공유기 외장하드의 전용 영역')]),
    '구조는 Linux 서버에서 필요한 자료를 전용 도구로 수집한 뒤, 암호화해서 SMB 공유에 보내는 방식입니다. 실제 배포에서는 공유기 외장하드를 목적지로 사용했습니다. 인증과 암호화 키는 서버의 보호된 경로에 두고 Git에는 넣지 않습니다. 현재 사례는 SMB 공유를 대상으로 구현되어 있습니다. 서버 전체의 부팅 가능한 디스크 복제와는 복구 목표가 다르다는 점도 설명서에 적었습니다.',50)
add('프로젝트마다 필요한 파일만 꺼냅니다','BACKUP',cards([
    ('project-alpha','코드 · Git · 환경 설정 → 별도 암호화 파일'),('project-beta','다른 프로젝트 → 별도 암호화 파일'),('구성요소','시스템 설정 · DB dump · 메일 export')]),
    '프로젝트를 하나의 거대한 파일에만 넣지 않고 디렉터리별로 나눴습니다. 그림의 프로젝트 이름은 만든 예시입니다. 시스템 설정과 데이터베이스 덤프도 별도 파일로 저장합니다. 모든 구성요소가 검증된 뒤 하루 성공 묶음 하나로 기록하며 필요한 프로젝트나 구성요소를 골라 복원할 수 있습니다. 파일이 나뉘어 있다는 것과 성공 기준이 나뉘는 것은 구분했습니다.',45)
add('검증이 끝나야 그날의 성공입니다','BACKUP',flow([
    ('연결 확인','정확한 CIFS mount와 표식'),('수집·암호화','DB·메일 전용 export'),('되읽기','전송본 해시와 복호화 검사'),('성공·보관','성공 5일 · 새 임시 수집본 정리')]),
    '전송이 끝났다는 것만으로 성공 처리하지 않습니다. 올바른 CIFS 공유인지 확인하고 수집한 뒤 암호화하여 전송합니다. 전송본을 다시 읽어 해시와 복호화를 확인한 다음에야 성공 기록을 만듭니다. 새 성공본이 검증된 뒤 승인된 범위의 오래된 성공본을 정리합니다. 공간이 없거나 실패하면 기존 성공본부터 지우지 않습니다. 실패 조각과 다른 프로그램의 자료도 자동 정리 범위에 넣지 않았습니다.',55)
add('오류 22를 암호 문제로 단정하지 않았습니다','CASE 01',
    '<div class="case-number diagram"><div><strong>22</strong><span>SMB 연결 옵션 오류</span></div><div class="arrow">→</div><div><strong>2.0</strong><span>이 공유기에서 확인한 연결 설정</span></div></div><div class="tag">실패 보고 → 협상 확인 → 원인에 맞게 수정 → 재시험</div>',
    '첫 실제 문제는 SMB 연결 오류였습니다. 처음에는 안내가 인증 문제처럼 보여 사용자가 암호를 반복 입력했습니다. 오류 번호와 협상 결과를 확인하니 이 공유기에서 필요한 버전 설정을 맞추지 못한 문제였습니다. 이 사례에서는 vers 2.0 설정으로 연결됐습니다. 모든 SMB 서버에 같은 버전을 권장한다는 뜻은 아닙니다. 같은 입력을 반복하기 전에 원인을 조사한다는 규칙이 실제로 필요한 순간이었습니다.',60)
add('“메일 중단 없이 진행하고 싶습니다”','CASE 02',
    '<div class="compare"><div><h3>요구 변경</h3><p>서비스 중단 없이 수집<br>기존 중단 경로 제거</p></div><div class="good"><h3>새 확인 기준</h3><p>온라인 export 검사<br>운영 상태 전후 비교</p></div></div>',
    '사용자는 메일을 중단하지 않고 싶다고 했습니다. 그러면 처음 만든 계획과 완료 기준도 바꿔야 합니다. 중단 동의를 계속 요청하는 대신 서비스 제어 경로를 제거하고 Dovecot, SQLite, Redis의 온라인 수집으로 나눴습니다. 실제 운영 메일을 중단하는 방식은 실행하지 않았습니다. 요구가 바뀌면 카드와 설명서, 시험도 함께 고쳐야 한다는 점을 보여 주는 사례입니다.',55)
add('중단 없는 수집의 한계도 남깁니다','CASE 02',cards([
    ('우편함','사용자별 Dovecot export · 격리 읽기'),('설정과 DB','SQLite backup · Redis RDB 검사'),('시점의 한계','전체 공통 시점 아님 · SMTP 큐 관측 시 실패')]),
    '중단 없이 수집했다고 해서 모든 자료가 완전히 같은 순간의 상태가 되는 것은 아닙니다. 우편함은 사용자별로 복사하고 SQLite와 Redis도 각각 전용 절차를 사용합니다. SMTP 대기열은 검사 시점에 비어 있어야 하며 메일이 관측되면 기존 서비스를 유지하고 백업을 실패로 남깁니다. 전체 공통 시점이나 모든 송수신 상태를 보장했다고 쓰지 않았습니다. 제한을 결과와 함께 적는 것이 복구 계획에 도움이 됩니다.',55)
add('합성 시험과 실제 복구 증거를 구분합니다','VERIFY',cards([
    ('123개 자동 시험','암호화 · 누락 · 손상 · 경로 · 승인 경계'),('실제 백업 읽기','설정 파일 · 메일 SQLite · 우편함 7개'),('실제 DB 복원','격리 PostgreSQL · DB 5개 · 읽기 검사')]),
    '숫자가 서로 다른 의미를 갖는다는 점을 강조하고 싶습니다. 123개는 합성 자료로 확인한 자동 시험의 개수입니다. 실제 공유기에 저장한 백업에서는 설정 파일과 메일 SQLite, 일곱 우편함의 읽기를 확인했습니다. 별도 PostgreSQL에서는 다섯 데이터베이스를 복원하고 읽었습니다. 이 결과를 전체 메일 서비스 로그인이나 애플리케이션 기능 검증으로 확대하지 않았습니다. 시험 이름보다 실제로 어디까지 실행했는지가 중요합니다.',55)
add('복원해 보고 용량도 측정했습니다','VERIFY',
    '<div class="metrics diagram"><div><strong>2.819<span>GiB</span></strong><p>첫 암호화 성공본</p></div><div><strong>5<span>DB</span></strong><p>격리 복원·읽기 확인</p></div><div><strong>7<span>우편함</span></strong><p>격리 읽기 확인</p></div></div><div class="tag">같은 크기 5개 ≈ 14.094GiB · 새 성공본 추가 순간 6개 ≈ 16.912GiB</div>',
    '첫 성공본의 암호화 크기는 약 2.819GiB였습니다. 같은 크기 다섯 개는 약 14.094GiB이고 새 성공본을 확인한 뒤 지우기 위해 잠시 여섯 개를 두면 약 16.912GiB입니다. 사용자가 정한 50GiB 범위 안이라는 현재 크기의 추정입니다. 앞으로의 증가나 실패 조각, 임시 공간까지 무한히 충분하다는 뜻은 아닙니다. 운영 데이터베이스의 원래 행 수 기준은 없었으므로 원본과 행 수가 같다고도 주장하지 않았습니다.',60)
add('서버 밖의 키로 실제 백업을 열었습니다','RECOVERY',flow([
    ('외부 보관','운영자가 암호 관리자에 저장'),('숨김 재입력','보관한 곳에서 다시 꺼내 입력'),('공유기 복호화','재입력한 키로 명세 열기'),('결과 기록','키 값 없이 성공 여부만 남기기')]),
    '서버가 고장 났는데 키도 그 서버에만 있으면 암호화 백업을 열지 못할 수 있습니다. 운영자가 외부에 기존 키를 보관하고 다시 꺼내 숨김 입력했습니다. 도구는 그 키로 실제 공유기 백업 명세를 복호화했습니다. 방송에서는 키를 확인하거나 입력하는 장면을 찍지 않습니다. 에이전트가 외부 암호 관리자 자체를 검사한 것은 아니고 운영자 확인과 실제 재입력 결과를 기록한 것입니다.',50)
add('스냅샷이 오류 재현 자료가 됐습니다','CASE 03','',
    '자동 실행을 설치할 때 또 다른 오류가 나왔습니다. 처음 연결할 때 만든 설정에는 나중에 추가한 형식 항목이 없었고, 백업기는 기본값을 적용했지만 설치기는 이를 누락 오류로 처리했습니다. 보존한 열 번째 단계의 설정 예제로 같은 거부를 재현했습니다. 설치기도 같은 기본값을 사용하도록 고쳤고 실제 재적용이 통과했습니다. 중간 스냅샷이 영상 장면뿐 아니라 이전 상태를 재현하는 근거가 된 장면입니다.',60,'step','17-config-compatibility')
add('사용자 의견: 승인 입력이 너무 길었습니다','CASE 04',
    '<div class="case-number diagram"><div><strong class="words">긴 승인 문장</strong><span>정책 전체를 다시 입력</span></div><div class="arrow">→</div><div><strong>y</strong><span>계획을 읽고 승인 · Enter는 취소</span></div></div><div class="tag">이미 저장된 승인 정책은 유지 · 재실행 요청 없음</div>',
    '설치가 성공한 뒤 사용자가 승인 문구가 너무 길다고 말했습니다. 앞으로는 계획을 읽은 뒤 y 한 글자로 승인하고 Enter는 취소하도록 바꿨습니다. 여기서 중요한 점은 이미 저장한 정책의 기준 문장까지 바꾸지 않았다는 것입니다. 그 값을 바꾸면 매일 실행기가 기존 승인을 거부할 수 있었습니다. 입력만 짧게 바꾸고 이전 승인 기록이 그대로 유효한지 시험했습니다. 이미 설치한 사용자에게 다시 입력을 요구하지 않았습니다.',45)
add('일일 예약 활성과 첫 실행 성공은 다릅니다','OPERATE','',
    '이 화면은 일일 예약과 공유 연결 설정이 실제로 활성화된 시점의 공개 진행 카드입니다. 다음 예약은 사례 당시 한국 시간10월21일 오전4시였습니다. timer가 켜져 있다는 사실과 그 예약이 실제 성공했다는 사실은 다릅니다. 자료 제작 기준으로 첫 예약 성공과 재부팅 연결, 메일 전체 서비스 복구는 남아 있습니다. 서버의 기본 Result 성공 표시를 실행 증거로 오해하지 않도록 실제 시작 시각도 확인했습니다.',50,'step','18-daily-enabled')
add('시연 1 · 공개 저장소와 대본을 엽니다','DEMO',screen('방송용 시연 · 공개 자료만',
    '<div class="checks"><p>1. README의 dev-process-light 적용 흐름</p><p>2. lecture.html에서 P = 발표 · N = 대본</p><p>3. 공개 진행 카드와 그 단계의 설명 보기</p></div>'),
    '첫 시연은 공개 자료만 사용합니다. 공개 저장소의 README에서 원본 dev-process-light 링크와 이번 사례의 작업 흐름을 보여 줍니다. 이어서 이 강의의 HTML 파일을 열고 P를 눌러 발표 모드, N을 눌러 대본을 보여 줍니다. 단계 스냅샷은 당시 만든 진행 카드라는 점을 설명합니다. 원본 운영 프로젝트의 조사 문서나 보호된 터미널로 전환하지 않습니다. 실제 화면 전환 순서는 별도 시연 대본에 준비해 두었습니다.',90,'step')
add('시연 2 · 합성 자료 시험을 실행합니다','DEMO',screen('공개 체크아웃에서 · 관리자 권한 없이',
    '<pre>python3 -m unittest discover -s tests -v</pre><p>합성 암호화·복원·승인 경계 시험<br>운영 백업·서비스 제어·실제 키 입력은 실행하지 않음</p>'),
    '두 번째 시연은 공개 체크아웃의 합성 자료 시험입니다. 관리자 권한 없이 표시된 명령을 실행합니다. 결과가 내려가는 동안 손상 파일, 잘못된 경로, 빈 승인 입력처럼 실패해야 하는 경우도 시험한다는 점을 설명합니다. 처음 녹화할 환경에서는 Python과 GnuPG가 준비됐는지 미리 확인합니다. 시간이 부족하면 동봉한 공개 시험 결과를 열어 볼 수 있지만 방금 실행한 것처럼 소개하지 않습니다. 운영 백업이나 실제 키 확인 명령은 시연하지 않습니다.',90,'step')
add('여러분의 작은 작업에도 적용해 보세요','TAKEAWAY',flow([
    ('원본 읽기','설치 안내와 버전 확인'),('카드 한 장','완료 기준을 세 가지 적기'),('작게 검증','합성 자료로 실패도 확인'),('함께 저장','기록·설명서·Git을 한 단계로')]),
    '백업 도구 전체를 그대로 도입할 필요는 없습니다. 먼저 원본 dev-process-light 저장소에서 설치 안내를 읽고 작은 작업 하나에 적용해 보세요. 완료 기준을 세 가지 정도로 적고, 합성 자료로 성공과 실패를 확인한 뒤, 설명서와 기록을 같은 Git 저장에 넣어 봅니다. 이 사례의 공개 코드는 특정 서버 구성을 일반 예시 이름으로 바꾼 학습 자료입니다. 실제 설치에는 자신의 경로와 서비스 구성을 확인하고 별도 검증해야 합니다.',50)
add('기록은 다음 복구와 다음 개발을 돕습니다','CLOSING',
    '<div class="closing"><p>요구 → 코드 → 실행 증거 → 다음 작업</p><div class="links">github.com/sungbeom78/dev-process-light<br>github.com/sungbeom78/smb-backup-lab</div><div class="tag">적용 버전 고정 · 확인 범위 공개 · 비밀 없는 자료</div></div>',
    '오늘 사례의 핵심은 요구가 코드로 이어지고 실행 증거가 다시 다음 작업으로 이어졌다는 점입니다. 기록은 영상을 만들 때도 도움이 됐지만 복구를 준비하고 오류를 재현할 때 더 직접적으로 쓰였습니다. 원본 dev-process-light와 smb-backup-lab 링크는 화면과 설명란에 넣었습니다. 확인하지 않은 부분은 계속 남은 시험으로 표시하겠습니다. 여러분의 작은 프로젝트에서도 이 흐름을 한 번 적용해 보시면 좋겠습니다.',35,'cover')
add('부록 · 자주 받는 질문','Q&A',
    '<table class="grid"><tr><th>질문</th><th>이번 사례의 답</th></tr><tr><td>다른 저장소도 지원하나요?</td><td>현재 구현 대상은 SMB 공유입니다.</td></tr><tr><td>서버 전체를 바로 부팅하나요?</td><td>OS 재설치 후 설정·데이터 복원이 목표입니다.</td></tr><tr><td>메일 전체 복구를 끝냈나요?</td><td>SQLite·우편함 읽기 확인, 전체 로그인/송수신은 별도입니다.</td></tr><tr><td>실패한 날 이전 성공본을 지우나요?</td><td>새 성공 검증이 없으면 보관 정리를 하지 않습니다.</td></tr></table>',
    '질의응답에서 자주 나올 수 있는 범위를 정리했습니다. 현재 대상은 SMB 공유이며 다른 저장소 지원을 구현했다고 소개하지 않습니다. 서버 전체를 부팅 가능하게 복제하는 백업도 아닙니다. 메일은 확인한 SQLite와 우편함 읽기, 남은 전체 서비스 검증을 구분합니다. 실패한 날에는 새 성공본 검증이 없으므로 기존 성공본 보관 정리가 진행되지 않습니다. 이 페이지는 본편 시간에서 제외한 선택 부록입니다.',0)

CSS='''
:root{--ink:#112c41;--muted:#537083;--bg:#dce4e9;--paper:#f8fafb;--soft:#e9f1f3;--accent:#103449;--accent-2:#0b877c;--line:#cad9df}
.slide>.s{gap:.7em}.slide .kicker{letter-spacing:.12em;font-size:.68em}.slide h2{font-size:1.94em;letter-spacing:-.04em}.slide[data-kind=cover]{background:#102d43}.slide[data-kind=cover] h1{font-size:2.8em;letter-spacing:-.05em;line-height:1.22}.slide[data-kind=cover] .kicker{color:#62d9c6}
.slide .cards{gap:.8em}.slide .card{border:0;background:#e8f0f2;border-top:.22em solid #168a7d;min-height:9em;padding:1.1em}.slide .card b{font-size:1.32em;line-height:1.25;margin-bottom:.7em}.slide .card span{font-size:1em;line-height:1.6}.tag{background:#e2efec;color:#16483f;border-radius:2em;padding:.4em .9em;font-size:.8em;align-self:flex-start}.hero{display:flex;flex-direction:column;gap:1.1em;color:#e6f2f6}.hero .eyebrow{color:#6be0cd;font-size:1.3em}.hero>.tag,.closing>.tag{background:#204a57;color:#9cedd8}
.checks{display:flex;flex-direction:column;gap:.55em;font-size:1.2em}.slide .flow{gap:.65em}.slide .flow li{padding:.9em .55em}.slide .flow li b{font-size:1.1em}.slide .flow li span{font-size:.85em;margin-top:.55em}.slide .body{gap:1em}.slide .grid{font-size:.9em}.slide .grid td,.slide .grid th{padding:.5em .7em}.timeline{display:grid;grid-template-columns:repeat(5,1fr);gap:.7em}.timeline>div{border-top:.25em solid #168a7d;background:#e8f0f2;display:flex;flex-direction:column;padding:1.2em .7em;gap:.9em;min-height:10em}.timeline b{font-size:1.5em;color:#0b877c}.timeline span{line-height:1.7;font-size:.9em}
.case-number{display:flex;align-items:center;justify-content:space-around;gap:1em}.case-number>div{display:flex;flex-direction:column;gap:.6em;align-items:center}.case-number strong{font-size:5em;line-height:1;color:#0b877c}.case-number .words{font-size:2em}.case-number .arrow{font-size:2em;color:#97aeb9}.case-number span{font-size:.9em}.metrics{display:grid;grid-template-columns:1.7fr 1fr 1.2fr;gap:1em}.metrics>div{background:#e8f0f2;padding:1.1em;border-top:.2em solid #168a7d}.metrics strong{font-size:2.8em;letter-spacing:-.05em;color:#0b877c}.metrics strong span{font-size:.35em;margin-left:.3em;letter-spacing:0}.metrics p{font-size:.9em;margin-top:.8em}.compare{display:flex;gap:1em}.compare>div{flex:1;background:#e8f0f2;padding:1.5em;border-left:.2em solid #f1aa53}.compare>.good{border-color:#168a7d}.compare h3{margin:0 0 .8em;font-size:1.25em}.compare p{line-height:1.7}
.closing{display:flex;flex-direction:column;gap:1.3em}.closing>p{font-size:1.5em;color:#fff}.closing .links{color:#6be0cd;font-size:1.1em;line-height:1.8}.slide .foot{position:absolute;bottom:.95em;left:3.7em;font-size:.62em;color:#688494}.slide[data-kind=cover] .foot{color:#8db8c5}.snapshot>.s{padding:0;gap:0}.snapshot>.s>.kicker,.snapshot>.s>h2,.snapshot>.s>.foot{display:none}.snapshot .body{padding:0;gap:0}.snapshot .body img{width:100%;height:100%;object-fit:contain}.snapshot::after{background:#102d43;color:white;padding:.1em .5em;border-radius:.2em}
'''

def write_texts():
    OUT.mkdir(parents=True,exist_ok=True)
    for name in ('assets','slides','previews'):(OUT/name).mkdir(exist_ok=True)
    for name in ('manual.css','manual.js'):
        shutil.copyfile(ROOT/'manual/assets'/name,OUT/'assets'/name)
    (OUT/'assets/broadcast.css').write_text(CSS)
    elapsed=0;sections=[];script=['# dev-process-light × smb-backup-lab — 전체 강의 대본\n',
        '약20분 본편 + 선택 Q&A. 시간은 권장값이며 실제 녹음 속도에 맞춰 조정합니다. 운영 비밀 화면은 시연하지 않습니다.\n']
    timing=['# 방송 진행표\n','|시작|슬라이드|주제|권장 시간|화면|','|---|---|---|---|---|']
    for i,s in enumerate(DECK,1):
        s['number']=i;s['start_seconds']=elapsed
        s['plain_title']=s['title'].replace('<br>',' ')
        body=s['body'];cls=''
        if s['image']:
            name=s['image']+'.png'
            shutil.copyfile(ROOT/'media/snapshots'/s['image']/'screen.png',OUT/'assets'/name)
            body=f'<img src="assets/{name}" alt="{html.escape(s["plain_title"])}">';cls=' snapshot'
        head='h1' if s['kind']=='cover' else 'h2'
        sections.append(f'<section class="slide{cls}" data-kind="{s["kind"]}" data-title="{html.escape(s["plain_title"])}"><p class="kicker">{s["chapter"]} · DEV-PROCESS-LIGHT</p><{head}>{s["title"]}</{head}><div class="body">{body}</div><p class="foot">smb-backup-lab · 제작 기준 2026.10.20 · 검증 범위는 대본에 명시</p><aside class="notes">{html.escape(s["notes"])}</aside></section>')
        start=f'{elapsed//60:02}:{elapsed%60:02}'
        script.append(f'## {i:02}. {s["plain_title"]} · {start} · {s["seconds"]}초\n\n{s["notes"]}\n')
        timing.append(f'|{start}|{i:02}|{s["plain_title"]}|{s["seconds"]}초|'+('공개 체크아웃 / 합성 시험' if s['chapter']=='DEMO' else '슬라이드')+'|')
        elapsed+=s['seconds']
    document='''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>dev-process-light × smb-backup-lab</title><link rel="stylesheet" href="assets/manual.css"><link rel="stylesheet" href="assets/broadcast.css"><script defer src="assets/manual.js"></script></head><body data-manual="dev-process-light 실전 강의" data-index="lecture.html">'''+''.join(sections)+'</body></html>'
    (OUT/'lecture.html').write_text(document)
    (OUT/'speaker-notes.md').write_text('\n'.join(script))
    (OUT/'broadcast-runbook.md').write_text('\n'.join(timing)+f'\n\n본편 권장 합계: {elapsed//60}분 {elapsed%60}초. 마지막 부록은 선택입니다.\n\n'+
        '녹화 순서: 표지 → 슬라이드01–21 → 공개 체크아웃 시연22–23 → 정리24–25 → 선택 Q&A26.\n\nOBS는 슬라이드/브라우저를1920×1080 장면으로 준비합니다. 알림과 개인 계정 메뉴를 닫고 공개 체크아웃만 연 터미널을 별도 창으로 준비합니다. 실제 서버 SSH 창·암호 관리자·/etc·백업 원문은 화면 전환 대상에서 제외합니다. 마이크 입력과 발표자 노트가 방송 화면에 섞이지 않는지10초 시험 녹화로 확인합니다.\n\nPPTX의 발표자 보기 또는 HTML의N키는 녹음자 참고용입니다. 대본을 보여 주는 기능을 소개하는 장면 외에는 노트를 숨깁니다. 진행표는 권장 시간이며 실제 녹음 길이에 맞춰 자막과 챕터 시간을 수정합니다.\n')
    (OUT/'demo-script.md').write_text('''# 시연 대본 — 공개 자료만 사용

## 1. 시작 전
공개 체크아웃은 /project/smb-backup-lab, 운영 원본은 별도입니다. 공개용 브라우저/터미널을 준비하고 탭 주소와 창 제목을 확인합니다. 발표 자료는 이 폴더의 lecture.html을 오프라인으로 열면 됩니다. P 발표, N 대본, 방향키 이동, Esc 돌아가기입니다.

## 2. README와 흐름 (슬라이드22, 약90초)
1. 공개 저장소 README의 **Built with dev-process-light** 부분을 보여 줍니다.
2. 원본 링크와 적용 버전, 시작→합의→확인→저장 그림을 설명합니다.
3. lecture.html에서 P/N을 누르고17단계 공개 카드를 보여 줍니다. 과거 시점에 만든 진행 카드임을 말합니다.

## 3. 합성 시험 (슬라이드23, 약90초)
관리자 권한 없이 공개 체크아웃에서 실행합니다. Python3와GnuPG가 필요하며 운영 DB/Docker/공유기 접속 정보는 필요 없습니다.

```bash
cd /project/smb-backup-lab
python3 -m unittest discover -s tests -v
```

말하기: “성공해야 하는 경우와 손상·잘못된 경로·승인 거부처럼 실패해야 하는 경우를 함께 검사합니다.” 실패가 나면 원인을 확인합니다. 동봉된검증 요약은 이전 실행 결과로 소개하며 방금 실행했다고 말하지 않습니다.

## 4. 종료
강의 화면24쪽으로 돌아와 작은 작업 카드부터 적용하도록 안내합니다. 마지막 링크를 설명란에도 제공합니다. 실제 설치/마운트/백업/키 확인 명령은 이번 방송에서 실행하지 않습니다. 테스트 중 만들어진 합성 임시 파일만 시험 도구가 정리합니다.
''')
    (OUT/'youtube-description.md').write_text(f'''# 유튜브 제목·설명란 초안

## 추천 제목
AI와 만드는 서버 백업 — dev-process-light 실전 적용기 | smb-backup-lab

## 설명란
AI와 프로젝트를 이어갈 때 필요한 작업 카드·검증·기록·Git 저장을 서버 백업 개발 사례로 소개합니다. SMB 연결 오류, 중단 없는 메일 수집, 실제 DB 복구 확인과 사용자 피드백까지 어떻게 기록했는지 보여 드립니다.

- 원본 dev-process-light: {UPSTREAM}
- 사례 코드와 방송 자료: {REPO}
- 적용 버전: {PIN}

제작 기준2026-10-20: 첫 공유기 암호화 백업, 설정 파일·메일SQLite·우편함7개 읽기, PostgreSQL5개 격리 복원, 외부 복구 키 재입력과 일일 예약 활성까지 확인했습니다. 첫 예약 실행 성공·재부팅 공유 연결·Mailu 전체 로그인/송수신·앱/보존 볼륨 전체 복구는 추가 시험입니다. 실제 계정·키·DB 행·원본 서버 화면은 자료에 넣지 않았습니다.

챕터 시작점은 broadcast-runbook.md의 권장값을 참고하되 실제 편집 영상의 길이에 맞춰 입력하세요. 아직 녹음/영상 업로드는 하지 않았습니다.

#devprocesslight #SMB #Linux #백업 #복구 #AI개발
''')
    (OUT/'README.md').write_text('''# dev-process-light × smb-backup-lab 방송 패키지

- `smb-backup-lab-dev-process-light.pptx`: 16:9 발표용. 26쪽, 발표자 노트에 전체 강의 대본.
- `smb-backup-lab-dev-process-light.pdf`: 같은 화면의 배포용 PDF.
- `lecture.html`: 오프라인 발표. P 발표, N 대본, ←/→ 이동.
- `speaker-notes.md`: 슬라이드별 전체 한국어 강의 대본.
- `broadcast-runbook.md`: 약20분 본편 시간표·녹화 장면/준비 순서.
- `demo-script.md`: 공개 자료와 합성 시험으로 진행하는 안전한 시연 대본.
- `youtube-description.md`: 제목·설명란·출처·해시태그 초안.
- `previews/contact-sheet.jpg`: 전체 슬라이드 미리보기.
- `slides/`: 1920×1080 PNG, 썸네일/편집 컷으로 사용 가능.
- `deck.json`, `manifest.json`: 내용/대본·검증/파일 해시.
- `smb-backup-lab-broadcast-kit.zip`: 위 자료를 한 번에 받는 묶음.

PPTX 화면은 글꼴과 배치를 보존한 이미지형 슬라이드이며 발표자 노트는 편집할 수 있습니다. 화면 글자를 바꾸려면 `scripts/build_broadcast.py`의 내용/도식을 수정해 재생성합니다. 오디오 녹음·실제 영상 완성본·유튜브 업로드는 포함하지 않습니다. 실제 터미널/암호 화면을 공개 카드로 오해하지 않도록 화면의 출처와 시점을 대본에 남겼습니다.

적용 원본: https://github.com/sungbeom78/dev-process-light (MIT, 원본 라이선스 보존). 이 강의는 원본 방법을 실제 사례에 적용한 설명이며 원본 저장소의 공식 강의라는 뜻은 아닙니다.
''')
    (OUT/'deck.json').write_text(json.dumps({'title':'dev-process-light × smb-backup-lab','upstream':UPSTREAM,'upstream_commit':PIN,'slides':DECK,'recommended_main_seconds':elapsed},ensure_ascii=False,indent=2)+'\n')
    # A lecture chapter in the operator manual links to the canonical assets.
    link='../media/broadcast/2026-10-20/'
    local=document.replace('href="assets/',f'href="{link}assets/').replace('src="assets/',f'src="{link}assets/').replace('data-index="lecture.html"','data-index="index.html"')
    (ROOT/'manual/lecture.html').write_text(local)


def render():
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image, ImageOps, ImageDraw
    from pypdf import PdfReader
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='dev-process-light × smb-backup-lab'
    prs.core_properties.author='sungbeom78'
    issues=[];images=[]
    with sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1920,'height':1080},device_scale_factor=1)
        page.goto((OUT/'lecture.html').as_uri()+'?export');page.evaluate('document.fonts.ready')
        page.add_style_tag(content='.slide{width:1920px!important;max-width:none!important;border-radius:0!important;box-shadow:none!important;margin:0!important}.m-deck{padding:0!important}.m-bar{display:none!important}')
        slides=page.locator('section.slide')
        if slides.count()!=len(DECK):raise RuntimeError('slide count mismatch')
        for i,s in enumerate(DECK):
            el=slides.nth(i)
            bad=el.evaluate('''el=>{const p=el.getBoundingClientRect();return Array.from(el.querySelectorAll('.body *')).filter(x=>{if(x.closest('aside'))return false;const r=x.getBoundingClientRect();return r.width&&r.height&&(r.left<p.left-1||r.right>p.right+1||r.top<p.top-1||r.bottom>p.bottom+1)}).map(x=>x.tagName)}''')
            if bad:issues.append({'slide':i+1,'overflow':bad})
            image=OUT/'slides'/f'{i+1:02}.png';el.screenshot(path=str(image));images.append(image)
            sl=prs.slides.add_slide(prs.slide_layouts[6])
            sl.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
            sl.notes_slide.notes_text_frame.text=f'{i+1:02}. {s["plain_title"]}\n권장 {s["seconds"]}초 · {s["chapter"]}\n\n{s["notes"]}'
        browser.close()
    if issues:raise RuntimeError(json.dumps(issues,ensure_ascii=False))
    outpptx=OUT/'smb-backup-lab-dev-process-light.pptx';prs.save(outpptx)
    rgb=[Image.open(p).convert('RGB') for p in images]
    pdf=OUT/'smb-backup-lab-dev-process-light.pdf'
    rgb[0].save(pdf,'PDF',resolution=144,save_all=True,append_images=rgb[1:])
    if len(PdfReader(pdf).pages)!=len(DECK):raise RuntimeError('PDF count mismatch')
    check=Presentation(outpptx)
    if len(check.slides)!=len(DECK) or any(s['notes'] not in sl.notes_slide.notes_text_frame.text for s,sl in zip(DECK,check.slides)):
        raise RuntimeError('PPTX notes mismatch')
    cols=4;w=480;h=270;gap=12;rows=(len(images)+cols-1)//cols
    sheet=Image.new('RGB',(cols*(w+gap)+gap,rows*(h+gap)+gap),'#dce4e9')
    for i,p in enumerate(images):sheet.paste(Image.open(p).convert('RGB').resize((w,h)),(gap+(i%cols)*(w+gap),gap+(i//cols)*(h+gap)))
    sheet.save(OUT/'previews/contact-sheet.jpg',quality=92)
    bundle=OUT/'smb-backup-lab-broadcast-kit.zip'
    files={str(p.relative_to(OUT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(OUT.rglob('*')) if p.is_file() and p.name not in ('manifest.json',bundle.name)}
    (OUT/'manifest.json').write_text(json.dumps({'generated_at':datetime.now(ZoneInfo('Asia/Seoul')).isoformat(),
        'slide_count':len(DECK),'pptx_notes_verified':True,'pdf_pages_verified':True,'overflow_checks':issues,
        'resolution':[1920,1080],'main_seconds':sum(s['seconds'] for s in DECK),'files_sha256':files},ensure_ascii=False,indent=2)+'\n')
    with zipfile.ZipFile(bundle,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file() and p!=bundle:z.write(p,p.relative_to(OUT))
    print(json.dumps({'directory':str(OUT),'slides':len(DECK),'main_seconds':sum(s['seconds'] for s in DECK),'pptx':str(outpptx),'pdf':str(pdf)},ensure_ascii=False))


if __name__=='__main__':
    write_texts()
    if '--html-only' not in sys.argv:render()
