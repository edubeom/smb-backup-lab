#!/usr/bin/env python3
"""Record public/synthetic demonstrations and assemble a captioned broadcast kit.

No privileged command, production backup, credential, or service operation is used.
Run with requirements-media.txt installed; recording tests use the public checkout.
"""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import argparse
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import zipfile

import imageio_ffmpeg
from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import sync_playwright
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = Path('/project/smb-backup-lab')
OUT = ROOT / 'media/demo/2026-10-20'
BROADCAST = ROOT / 'media/broadcast/2026-10-20'
FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
FONT = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
IDENTITY = json.loads((ROOT/'config/project-identity.json').read_text())
URL = IDENTITY['repository']
CMD = ['python3', '-m', 'unittest', 'discover', '-s', 'tests', '-v']


def now():
    return datetime.now(ZoneInfo('Asia/Seoul')).isoformat(timespec='seconds')


def sha(path):
    return hashlib.file_digest(open(path, 'rb'), 'sha256').hexdigest()


def put_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def call(args):
    result = subprocess.run([str(x) for x in args], capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(result.stderr[-3500:])
    return result


STYLE = '''*{box-sizing:border-box}html,body{margin:0;width:1600px;height:900px;overflow:hidden}
body{background:#112d43;color:#f1f7fa;font-family:"Noto Sans CJK KR",sans-serif;padding:72px 84px}
.eyebrow{font-size:25px;letter-spacing:3px;color:#6ee0cc}h1{font-size:64px;line-height:1.27;margin:35px 0 42px;letter-spacing:-2px}
.rows{display:flex;flex-direction:column;gap:22px}.row{background:#1a3d52;border-left:6px solid #5dd7bd;border-radius:7px;padding:22px 28px;font-size:36px;line-height:1.5}
.row b{color:#76e7d0}.small{font-size:27px;color:#bed1dc}footer{position:absolute;bottom:48px;font-size:23px;color:#a7c4d4}
code{font-family:"Noto Sans Mono CJK KR",monospace;font-size:30px;overflow-wrap:anywhere}.result{font-size:110px;color:#76e7d0}
'''


def card(browser, name, title, rows, footer='공개 자료 · 방송 설명용 구성', eyebrow='DEV-PROCESS-LIGHT × SMB-BACKUP-LAB'):
    content = '<div class="rows">' + ''.join('<div class="row">' + r + '</div>' for r in rows) + '</div>'
    doc = f'<!doctype html><html lang="ko"><meta charset="utf-8"><style>{STYLE}</style><div class="eyebrow">{eyebrow}</div><h1>{title}</h1>{content}<footer>{html.escape(footer)}</footer></html>'
    path = OUT / 'sources' / f'{name}.html'
    path.write_text(doc)
    page = browser.new_page(viewport={'width': 1600, 'height': 900})
    page.goto(path.as_uri())
    page.evaluate('document.fonts.ready')
    # Fixed canvas intentionally hides overflow; inspect every visible element.
    overflowing = page.locator('h1,.row,footer').evaluate_all('es=>es.some(e=>{let r=e.getBoundingClientRect();return r.bottom>900||r.right>1600||r.x<0})')
    if overflowing:
        raise RuntimeError(f'card overflow: {name}')
    page.screenshot(path=str(path.with_suffix('.png')))
    page.close()
    return path.with_suffix('.png')


def save_recording(context, page, name):
    video = page.video
    context.close()
    video.save_as(str(OUT / 'sources' / f'{name}.webm'))
    video.delete()


def capture(public_only=False, tests_only=False):
    (OUT / 'sources').mkdir(parents=True, exist_ok=True)
    commit = call(['git', '-C', PUBLIC, 'rev-parse', 'HEAD']).stdout.strip()
    if call(['git', '-C', PUBLIC, 'status', '--porcelain']).stdout.strip():
        raise RuntimeError('Public checkout must be clean before recording')
    provenance = {'captured_at': now(), 'public_url': URL, 'public_commit': commit,
                  'production_access': False, 'source_kinds': ['actual public browser recording', 'actual synthetic test output viewer', 'historical progress card', 'editorial explanation card']}
    if public_only or tests_only:
        provenance=json.loads((OUT/'sources/provenance.json').read_text())
        provenance['public_views_recaptured_at' if public_only else 'tests_recaptured_at']=now()
        provenance['public_url']=URL
        provenance['public_commit']=commit
    with tempfile.TemporaryDirectory(prefix='smb-demo-record-') as td, sync_playwright() as p:
        browser = p.chromium.launch()
        def context():
            return browser.new_context(viewport={'width': 1600, 'height': 900}, record_video_dir=td,
                                       record_video_size={'width': 1600, 'height': 900})
        if not tests_only:
            ctx = context(); page = ctx.new_page(); started = time.monotonic()
            page.goto(URL, wait_until='domcontentloaded', timeout=60000)
            readme = page.locator('article.markdown-body')
            readme.wait_for()
            page.wait_for_timeout(3000)
            page.add_style_tag(content='html{scroll-behavior:auto!important}body{zoom:1.25}')
            page.wait_for_timeout(500)
            readme.locator('h1').first.evaluate('e=>e.scrollIntoView({block:"start",behavior:"instant"})')
            page.evaluate('window.scrollBy({top:-65,behavior:"instant"})')
            page.wait_for_timeout(1000)
            top=readme.locator('h1').first.evaluate('e=>e.getBoundingClientRect().top')
            if not 0 <= top < 180:raise RuntimeError(f'README heading outside recording frame: {top}')
            provenance['github'] = {'url': page.url, 'title': page.title(), 'trim_start': time.monotonic() - started,
                                    'adaptation': 'browser zoom 125%; actual public page, anonymous context'}
            page.screenshot(path=str(OUT / 'sources/github-readme.png'))
            page.wait_for_timeout(4500)
            page.mouse.wheel(0, 370); page.wait_for_timeout(4500)
            page.screenshot(path=str(OUT / 'sources/github-flow.png'))
            page.mouse.wheel(0, 250); page.wait_for_timeout(5000)
            save_recording(ctx, page, 'github-public')
            print('Captured public GitHub README', flush=True)

            ctx = context(); page = ctx.new_page(); started = time.monotonic()
            lecture = PUBLIC / 'media/broadcast/2026-10-20/lecture.html'
            page.goto(lecture.as_uri() + '#4')
            page.wait_for_timeout(500)
            provenance['lecture'] = {'source': 'media/broadcast/2026-10-20/lecture.html', 'sha256': sha(lecture),
                                     'trim_start': time.monotonic() - started, 'keys': ['P', 'ArrowRight', 'N', 'Escape', 'N']}
            page.wait_for_timeout(2000); page.keyboard.press('p'); page.wait_for_timeout(4500)
            assert page.locator('body').evaluate('e=>e.classList.contains("present")')
            page.screenshot(path=str(OUT / 'sources/lecture-present.png'))
            page.keyboard.press('ArrowRight'); page.wait_for_timeout(3000)
            page.keyboard.press('n'); page.keyboard.press('Escape')
            page.add_style_tag(content='body.show-notes .slide>aside.notes{font-size:24px;line-height:1.65;padding:25px}body.show-notes .slide .foot{display:none}')
            page.locator('section.slide').nth(4).locator('aside.notes').scroll_into_view_if_needed()
            provenance['lecture']['adaptation']='Notes viewing font enlarged to 24px; slide footer hidden while notes are visible; original notes text unchanged'
            page.wait_for_timeout(5000)
            assert page.locator('body').evaluate('e=>e.classList.contains("show-notes")')
            page.screenshot(path=str(OUT / 'sources/lecture-notes.png'))
            page.keyboard.press('n'); page.wait_for_timeout(3500)
            save_recording(ctx, page, 'lecture-keys')
            print('Captured actual lecture P/N keys', flush=True)
            if public_only:
                browser.close()
                put_json(OUT/'sources/provenance.json',provenance)
                return

        ctx = context(); page = ctx.new_page(); started = time.monotonic()
        viewer = OUT / 'sources/test-viewer.html'
        viewer.write_text('''<!doctype html><html lang="ko"><meta charset="utf-8"><style>
        *{box-sizing:border-box}body{margin:0;background:#0d2030;color:#e3edf3;font-family:"Noto Sans CJK KR",sans-serif;padding:34px 44px}
        h1{font-size:30px;color:#71e1c8;margin:0 0 12px}p{font-size:22px;color:#a8c2cf}code{font-size:27px}
        pre{height:610px;overflow:hidden;white-space:pre-wrap;overflow-wrap:anywhere;font-size:20px;line-height:1.35;background:#142e40;padding:20px;border:1px solid #31556c}
        #status{color:#ffd494;font-size:25px}</style><h1>이번 촬영 · 합성 시험의 실제 출력</h1>
        <p>공개 체크아웃 · 관리자 권한 없음 · 출력 뷰어</p><code>$ python3 -m unittest discover -s tests -v</code>
        <pre id="log"></pre><div id="status">실행 준비</div></html>''')
        page.goto(viewer.as_uri()); page.wait_for_timeout(500)
        provenance['tests'] = {'command': CMD, 'cwd_kind': 'reviewed public checkout', 'started_at': now(),
                               'trim_start': time.monotonic() - started, 'viewer': 'controlled browser output viewer, not desktop terminal'}
        page.wait_for_timeout(1000)
        log_path = OUT / 'sources/synthetic-tests.txt'
        page.locator('#status').evaluate('e=>e.textContent="시험 실행 중 · 종료 결과를 기다립니다"')
        run_start = time.monotonic()
        # Output stays in this source file and is displayed only after a privacy check.
        with log_path.open('w') as transcript:
            proc = subprocess.Popen(CMD, cwd=PUBLIC, env={**os.environ, 'PYTHONUNBUFFERED': '1'}, stdout=transcript, stderr=subprocess.STDOUT)
            while proc.poll() is None:
                text = log_path.read_text()
                check_text(text)
                page.evaluate('t=>{let e=document.querySelector("#log");e.textContent=t;e.scrollTop=e.scrollHeight}', text)
                page.wait_for_timeout(200)
        text = log_path.read_text(); check_text(text)
        result = re.search(r'Ran (\d+) tests? in ([\d.]+)s', text)
        if proc.returncode != 0 or not result or not text.rstrip().endswith('OK'):
            raise RuntimeError('Synthetic tests failed; recording will not show success')
        provenance['tests'].update(exit_code=proc.returncode, count=int(result[1]), reported_seconds=float(result[2]),
                                   elapsed_seconds=round(time.monotonic()-run_start, 3), finished_at=now(), transcript_sha256=sha(log_path))
        page.evaluate('t=>{let e=document.querySelector("#log");e.textContent=t;e.scrollTop=e.scrollHeight;document.querySelector("#status").textContent="종료 코드 0 · 실제 실행 결과 OK"}', text)
        page.screenshot(path=str(OUT / 'sources/tests-complete.png'))
        remaining = max(4500, int((23 - (time.monotonic()-started))*1000))
        page.wait_for_timeout(remaining)
        save_recording(ctx, page, 'synthetic-tests-live')
        print(f'Recorded {result[1]} synthetic tests: OK ({result[2]}s)', flush=True)

        card(browser, '01-intro', '시연 1<br>기록에서 근거를 찾기', ['<b>공개 저장소</b>에서 개발 흐름을 읽습니다.', '<b>강의 화면</b>과 당시의 진행 카드를 연결합니다.'])
        case = PUBLIC / 'docs/case-study.md'
        shutil.copyfile(case, OUT / 'sources/case-study.md')
        provenance['case_study'] = {'source': 'docs/case-study.md', 'sha256': sha(case), 'adaptation': 'excerpt paraphrased and rearranged for readability'}
        card(browser, '03-case', '작업 방법이 수정 근거가 됩니다', ['<b>요구</b> 매일 한 번 · 성공 5일 · 메일 중단 없음', '<b>확인</b> 오류 재현 → 수정 → 합성 시험', '<b>기록</b> 작업 카드 · 결과 · 다음 할 일'], '출처: 공개 docs/case-study.md · 읽기 쉽게 요약·재배치')
        shutil.copyfile(BROADCAST / 'assets/17-config-compatibility.png', OUT / 'sources/05-historical.png')
        historical = json.loads((ROOT / 'media/snapshots/17-config-compatibility/manifest.json').read_text())
        provenance['historical_card'] = {'source': 'media/snapshots/17-config-compatibility/screen.png',
                                         'captured_at': historical['captured_at'], 'kind': historical['screen_type'],
                                         'sha256': sha(OUT / 'sources/05-historical.png')}
        card(browser, '06-fix', '보관된 설정으로 문제를 재현', ['초기 설정에 <b>archive_layout</b> 항목이 없었습니다.', '백업기의 기본값을 설치기에도 적용했습니다.', '과거 설정 예제로 <b>호환 시험</b>을 남겼습니다.'], '출처: 공개 docs/case-study.md · 설명용 요약')
        card(browser, '07-transition', '다음은 실제 시험 화면입니다', ['기록을 읽는 데서 <b>직접 실행</b>으로 이어갑니다.', '합성 자료로 성공과 <b>실패할 조건</b>을 확인합니다.', '실행 결과와 운영 복구의 범위를 나눠 적습니다.'])
        card(browser, '08-intro', '시연 2<br>합성 자료로 직접 확인', ['이번 촬영에서 공개 코드의 시험을 실행합니다.', '결과는 <b>원본 출력과 실행 시각</b>으로 보존합니다.'])
        card(browser, '09-command', '공개 코드에서 시험을 실행', ['<code>python3 -m unittest discover -s tests -v</code>', '운영 백업 대신 <b>합성 파일·시험 설정</b>을 사용합니다.', '관리자 권한 없이 실행합니다.'])
        n = provenance['tests']['count']; seconds = provenance['tests']['reported_seconds']
        card(browser, '11-result', f'<span class="result">{n}</span>개 시험 통과', [f'실제 출력: <b>Ran {n} tests in {seconds:.3f}s</b>', '종료 코드 <b>0</b> · 최종 결과 <b>OK</b>', '이번 촬영의 원본 출력 파일을 함께 제공합니다.'], f'실행 완료: {provenance["tests"]["finished_at"]}')
        names = ['test_daily_failed_transfer_never_calls_prune', 'test_empty_or_negative_reply_never_installs_or_stores_approval', 'test_historical_setup_config_uses_transfer_default_without_mutation']
        for name in names:
            if not re.search(re.escape(name) + r'.*?\.\.\. ok', text):
                raise RuntimeError('Expected passing safety test absent: '+name)
        put_json(OUT / 'sources/selected-tests.json', {'source': 'synthetic-tests.txt', 'passed_test_names': names})
        card(browser, '12-safety', '“실패해야 하는 경우”도 시험', ['전송 실패 → <b>기존 성공본을 정리하지 않음</b>', '빈 승인 입력 → <b>설치와 승인 저장 없음</b>', '초기 설정 → <b>기본값을 적용하고 원본 유지</b>'], '출처: 이번 실제 시험 출력의 통과 항목 3개 · 전체 이름은 selected-tests.json')
        card(browser, '13-scope', '확인한 범위를 구분해서 말하기', ['<b>이번 촬영</b> 공개 코드의 합성 시험', '<b>과거 사례 기록</b> 첫 백업 · 격리 DB 복원 · 우편함 읽기', '<b>남은 확인</b> 첫 예약 성공 · 전체 메일/앱 복구'], '사례 상태 기준: 2026-10-20 공개 문서 · 운영 서비스를 새로 시험하지 않음')
        card(browser, '14-outro', '작게 만들고, 근거를 남깁니다', ['<b>시작 → 합의 → 확인 → 저장</b>', 'dev-process-light를 작은 작업 하나에 적용해 보세요.', '<span class="small">dev-process-light 원본 · 설명란 링크<br>github.com/edubeom/smb-backup-lab · hello@edubeom.com</span>'])
        browser.close()
    put_json(OUT / 'sources/provenance.json', provenance)
    print('Capture and explanation cards saved', flush=True)


def check_text(text):
    # Only public/synthetic output is eligible for the recording.
    banned = ['192.168.50.', 'lab-server', '/home/operator', '/etc/g1-backup/archive-passphrase',
              '-----BEGIN PRIVATE KEY-----']
    if any(s in text for s in banned) or re.search(r'(?:gh[pousr]_|github_pat_)[A-Za-z0-9_]{20,}', text):
        raise RuntimeError('Private identifier/token in media source; refusing to display')


def refresh_identity():
    """Refresh only editorial identity art; preserve original test recordings."""
    with sync_playwright() as p:
        browser=p.chromium.launch()
        card(browser, '14-outro', '작게 만들고, 근거를 남깁니다',
             ['<b>시작 → 합의 → 확인 → 저장</b>', 'dev-process-light를 작은 작업 하나에 적용해 보세요.',
              '<span class="small">dev-process-light 원본 · 설명란 링크<br>'+html.escape(URL.removeprefix('https://'))+'<br>'+html.escape(IDENTITY['contact_email'])+'</span>'])
        browser.close()
    path=OUT/'sources/provenance.json'
    provenance=json.loads(path.read_text())
    provenance['identity_update']={'updated_at':now(),'owner':IDENTITY['owner'],'contact_email':IDENTITY['contact_email'],
                                   'repository':URL,'original_test_recording_preserved':True,
                                   'framework_link_note':'Actual source repository retained; source has not moved.'}
    put_json(path,provenance)
    print('Identity card refreshed; original synthetic test evidence preserved',flush=True)


SCENES = [
    ('01-intro.png', 6, '시연 1 · 기록에서 근거 찾기', '설명용 구성', ['공개 자료에서 개발 흐름과 수정 근거를 찾아보겠습니다.']),
    ('github-public.webm', 14, '공개 README · 개발 흐름', '이번 촬영 · 실제 브라우저', ['README에서 Built with dev-process-light를 확인합니다.', '원본 도구와 이번 사례의 적용 범위를 연결합니다.']),
    ('03-case.png', 12, '작업 카드 → 확인 → 기록', '공개 문서 요약', ['요구와 완료 기준이 수정과 시험의 기준이 됩니다.', '결과와 다음 할 일을 같은 흐름으로 남깁니다.']),
    ('lecture-keys.webm', 18, 'HTML 강의 · P = 발표 / N = 대본', '이번 촬영 · 대본 보기 확대', ['P를 눌러 발표 화면으로 전환합니다.', 'N으로 대본을 켜고 Esc로 읽기 화면에서 확인합니다.', '발표 자료와 말할 내용을 함께 보관할 수 있습니다.']),
    ('05-historical.png', 14, '당시 진행 카드 · 설정 호환 수정', '과거 기록 · 서버 화면 아님', ['이 화면은 당시 제작한 공개용 진행 카드입니다.', '보관한 설정 예제로 설치기의 호환 오류를 재현했습니다.']),
    ('06-fix.png', 10, '오류 재현의 근거를 남기기', '공개 사례 설명', ['누락 항목의 기본값 처리를 백업기와 설치기에 맞췄습니다.', '같은 문제가 다시 생기지 않도록 시험을 추가했습니다.']),
    ('07-transition.png', 16, '기록에서 실제 실행으로', '다음 시연 안내', ['다음 장면에서는 공개 코드의 합성 시험을 실행합니다.', '성공뿐 아니라 거부해야 할 조건도 확인합니다.', '시험 결과와 운영 서비스 복구의 범위를 구분합니다.']),
    ('08-intro.png', 6, '시연 2 · 합성 자료로 직접 확인', '설명용 구성', ['이번 촬영에서 실행한 시험과 원본 출력을 보여 드립니다.']),
    ('09-command.png', 8, '공개 체크아웃의 시험 명령', '합성 자료 · 관리자 권한 없음', ['표시된 명령으로 공개 코드의 시험을 실행합니다.', '실제 운영 백업과 키 입력은 이 시연에 포함하지 않습니다.']),
    ('synthetic-tests-live.webm', 22, '합성 시험 실행 · 출력 보기', '이번 실행 · 원본 속도 발췌', ['이 장면은 이번 실행의 출력을 보여 주는 뷰어입니다.', '합성 파일로 암호화와 복원, 실패 조건을 시험합니다.', '화면에 나타난 결과의 원본도 함께 보존합니다.', '종료 코드와 전체 시험의 최종 결과를 확인합니다.']),
    ('11-result.png', 11, '실제 실행 결과 · 종료 코드 0', '이번 촬영의 결과', ['모든 시험이 통과했고 최종 출력은 OK입니다.', '시험 수와 실행 시간은 이번 원본 출력에서 가져왔습니다.']),
    ('12-safety.png', 15, '통과한 안전장치 시험 3개', '실제 출력에서 선택', ['전송 실패 때 이전 성공본을 정리하지 않는지 확인합니다.', '빈 승인 입력은 설치와 승인 저장을 막습니다.', '초기 설정과의 호환성도 계속 시험합니다.']),
    ('13-scope.png', 14, '합성 시험과 운영 검증의 범위', '공개 사례 상태 기준', ['이번 합성 시험을 운영 전체 복구 성공으로 소개하지 않습니다.', '과거 복구 기록과 남은 서비스 시험을 구분해서 적습니다.']),
    ('14-outro.png', 14, '시작 → 합의 → 확인 → 저장', '강의 24쪽으로 연결', ['dev-process-light를 작은 작업 하나에 적용해 보세요.', '기록과 실행 근거가 다음 개발과 복구를 돕습니다.'])
]


def stamp(t, ass=False):
    hours, t = divmod(t, 3600); minutes, seconds = divmod(t, 60)
    return f'{int(hours)}:{int(minutes):02}:{seconds:05.2f}' if ass else f'{int(hours):02}:{int(minutes):02}:{int(seconds):02},{round((seconds%1)*1000):03}'


def subtitles(scenes, stem):
    events=[]; start=0
    for s in scenes:
        for i, text in enumerate(s['captions']):
            a=start+i*s['duration']/len(s['captions']); b=start+(i+1)*s['duration']/len(s['captions'])
            events.append((a,b,text))
        start += s['duration']
    (OUT / f'{stem}.ko.srt').write_text('\n'.join(f'{i}\n{stamp(a)} --> {stamp(b)}\n{t}\n' for i,(a,b,t) in enumerate(events,1)))
    header='''[Script Info]
ScriptType: v4.00+
PlayResX: 1920
PlayResY: 1080
WrapStyle: 0
[V4+ Styles]
Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding
Style: Default,Noto Sans CJK KR,42,&H00F5F8FA,&H000000FF,&H00102030,&H80102030,0,0,0,0,100,100,0,0,1,1,0,2,60,60,37,1
[Events]
Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text
'''
    (OUT / f'{stem}.ko.ass').write_text(header+''.join(f'Dialogue: 0,{stamp(a,True)},{stamp(b,True)},Default,,0,0,0,,{t}\n' for a,b,t in events))


def label_image(scene, path):
    im=Image.new('RGBA',(1920,1080),(0,0,0,0)); d=ImageDraw.Draw(im)
    d.text((55,18),scene['title'],font=ImageFont.truetype(FONT,34,index=1),fill='#eff6fa')
    font=ImageFont.truetype(FONT,23,index=1); badge=scene['kind']
    width=d.textlength(badge,font=font)
    d.rounded_rectangle((1854-width-25,17,1875,65),radius=18,fill='#204c56')
    d.text((1854-width-13,23),badge,font=font,fill='#81e3cf')
    d.line((55,79,1865,79),fill='#315567',width=2)
    im.save(path)


def render():
    prov=json.loads((OUT/'sources/provenance.json').read_text())
    scenes=[]; cursor=0
    for i,(source,duration,title,kind,captions) in enumerate(SCENES,1):
        s=dict(number=i,source='sources/'+source,duration=duration,start=cursor,title=title,kind=kind,captions=captions)
        s['trim_start']=prov.get({'github-public.webm':'github','lecture-keys.webm':'lecture','synthetic-tests-live.webm':'tests'}.get(source,''),{}).get('trim_start',0)
        scenes.append(s);cursor+=duration
    assert sum(s['duration'] for s in scenes[:7])==90 and cursor==180
    put_json(OUT/'timeline.json',{'resolution':[1920,1080],'fps':24,'duration':180,'audio':'silent AAC; no voice/music','scenes':scenes})
    (OUT/'overlays').mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='smb-demo-edit-') as td:
        td=Path(td);clips=[]
        for s in scenes:
            src=OUT/s['source']; overlay=OUT/'overlays'/f'{s["number"]:02}.png'; label_image(s,overlay)
            dest=td/f'{s["number"]:02}.mp4'
            args=[FFMPEG,'-y','-loglevel','error','-threads','2']
            if src.suffix=='.png': args+=['-loop','1','-framerate','24']
            else: args+=['-ss',str(s['trim_start'])]
            args+=['-i',src,'-i',overlay]
            vf='[0:v]scale=1458:820:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:98:color=0x0d2233,setsar=1,fps=24,tpad=stop_mode=clone:stop_duration=24[base];[base][1:v]overlay=0:0,format=yuv420p,fade=t=in:st=0:d=0.22,fade=t=out:st='+str(s['duration']-.22)+':d=0.22[v]'
            args+=['-filter_complex_threads','1','-filter_complex',vf,'-map','[v]','-t',str(s['duration']),'-c:v','libx264','-threads','2','-preset','veryfast','-crf','21','-an',dest]
            call(args);clips.append(dest)
            print(f'Encoded scene {s["number"]:02}/14: {s["duration"]}s',flush=True)
        for name, selected, subset in [('demo-01-records',clips[:7],scenes[:7]),('demo-02-tests',clips[7:],scenes[7:])]:
            subtitles(subset,name)
            listing=td/f'{name}.txt';listing.write_text(''.join(f"file '{x}'\n" for x in selected))
            ass=OUT/f'{name}.ko.ass'
            call([FFMPEG,'-y','-loglevel','error','-f','concat','-safe','0','-i',listing,'-f','lavfi','-i','anullsrc=r=48000:cl=stereo',
                  '-vf',f'subtitles=filename={ass}', '-t','90','-c:v','libx264','-threads','2','-preset','veryfast','-crf','21','-pix_fmt','yuv420p',
                  '-c:a','aac','-b:a','96k','-movflags','+faststart',OUT/f'{name}.mp4'])
            print(f'Captioned {name}: 90s',flush=True)
        subtitles(scenes,'demo-combined')
        listing=td/'combined.txt';listing.write_text(''.join(f"file '{OUT/name}'\n" for name in ['demo-01-records.mp4','demo-02-tests.mp4']))
        # Video is already captioned. Stream-copy exact 90-second H.264 parts.
        call([FFMPEG,'-y','-loglevel','error','-f','concat','-safe','0','-i',listing,'-c','copy','-movflags','+faststart',OUT/'demo-combined.mp4'])
    print('Rendered two parts and 180-second combined video',flush=True)


def package():
    previews=OUT/'previews';previews.mkdir(exist_ok=True)
    for name in ['demo-01-records','demo-02-tests']:
        call([FFMPEG,'-y','-loglevel','error','-ss','2','-i',OUT/f'{name}.mp4','-frames:v','1',previews/f'{name}.png'])
    prs=Presentation(BROADCAST/'smb-backup-lab-dev-process-light.pptx')
    for i,name in [(21,'demo-01-records'),(22,'demo-02-tests')]:
        prs.slides[i].shapes.add_movie(str(OUT/f'{name}.mp4'),0,0,prs.slide_width,prs.slide_height,
                                     poster_frame_image=str(previews/f'{name}.png'),mime_type='video/mp4')
        prs.slides[i].notes_slide.notes_text_frame.text += '\n\n삽입 영상: '+name+'.mp4 · 90초 · 클릭하여 재생 · 무음/자막 포함. 이 영상이 해당 시연을 대체하므로 원래 90초 시연을 추가하지 않습니다.'
    prs.save(OUT/'smb-backup-lab-demo-embedded.pptx')
    timeline=json.loads((OUT/'timeline.json').read_text())
    rows=['# 시연 편집표와 직접 읽는 내레이션\n','영상은 무음이며 한국어 자막이 화면에 포함됩니다. 각 시연은 강의22·23쪽의 원래90초를 대체합니다. 합본은3분입니다.\n',
          '|합본 시각|시연 내부|장면|설명할 내용|','|---|---|---|---|']
    for s in timeline['scenes']:
        t=s['start']; local=t if t<90 else t-90
        rows.append(f'|{t//60:02}:{t%60:02}–{(t+s["duration"])//60:02}:{(t+s["duration"])%60:02}|시연{1 if t<90 else 2} · {local//60:02}:{local%60:02}|{s["title"]}|'+ ' '.join(s['captions'])+'|')
    (OUT/'narration-and-edit-cues.md').write_text('\n'.join(rows)+'\n\n자막 문장을 바탕으로 각 구간에 맞춰 직접 설명하세요. 실제 출력 뷰어는 터미널 화면을 흉내 낸 성공 연출이 아니라 같은 시각 실행한 표준 출력의 보기입니다. 과거 진행 카드는 실제 서버 녹화가 아닙니다.\n')
    (OUT/'README.md').write_text('''# 강의 시연 영상 묶음

작성자: edubeom · 연락처: hello@edubeom.com

dev-process-light × smb-backup-lab. 1920×1080, 24fps, H.264/AAC MP4. 음성·배경음악 없이 한국어 자막이 화면에 들어 있습니다. 조용한 AAC 트랙은 재생 호환을 위한 것입니다.

1. `demo-01-records.mp4` — 90초. 공개 GitHub/실제 HTML P·N 조작/과거 오류 해결 카드.
2. `demo-02-tests.mp4` — 90초. 이번에 실제 실행한 합성 시험과 결과/검증 범위.
3. `demo-combined.mp4` — 두 시연을 이어 붙인3분 합본.
4. `smb-backup-lab-demo-embedded.pptx` — 기존26쪽의22·23쪽에 각 영상을 포함한 사본. 원본 PPTX는 그대로입니다.
5. `narration-and-edit-cues.md` — 구간별 말할 내용. `timeline.json`은 초 단위 편집표.
6. `*.ko.srt`/`*.ko.ass` — 편집용 자막. MP4에는 이미 자막이 있으므로 그대로 방송할 때 자막을 중복 적용하지 마세요.
7. `sources/` — 원본 WebM 촬영, 공개 출처 사본, 실제 시험 출력/시각/공개commit, 설명 카드 원본.
8. `previews/` — 첫 장면과 주요 프레임 미리보기. `manifest.json` — 파일 해시와 검사 결과.

## 강의에 넣는 순서

- 기존 강의21쪽까지 설명한 뒤22쪽에서 시연1 영상을 클릭해 재생합니다. 영상이 끝나면23쪽 시연2를 재생하고24쪽으로 넘어갑니다. 각 영상이 원래 시연90초를 대체하므로 강의 권장 길이는21분55초로 같습니다.
- 영상 포함 PPTX는 ZIP 전체를 풀고 개인 PC의 PowerPoint에서 열어22·23쪽을 클릭해 보세요. 미디어는 PPTX 안에 포함했습니다. 여기서는 슬라이드/영상 관계·내장 파일·전체 디코딩까지 검사했으며 실제 PowerPoint 앱 재생은 실행 환경이 없어 확인하지 못했습니다.
- PPTX에서 재생이 지원되지 않으면 OBS의 미디어 소스로 두 MP4를 각각 추가하면 됩니다. 1920×1080 장면에 맞춰 놓고 '소스가 활성화되면 재생 다시 시작'을 켠 뒤, 마이크로 직접 설명하세요.
- 시작 전에 시연1의 GitHub 장면과 시연2의 결과 장면을10초 시험 녹화해 글자와 마이크를 확인하세요. 영상에는 소리가 없으므로 발표자의 마이크를 사용합니다.

## 출처와 확인 범위

GitHub/HTML 키 조작과 합성 시험 출력은 이번 촬영입니다. 브라우저는 익명 새 컨텍스트로 열었고 GitHub는125% 확대했습니다. 시험 화면은 실제 실행의 표준 출력을 읽는 별도 보기이며 전체 원본은 synthetic-tests.txt에 있습니다. 영상에는 원래 속도의 발췌가 들어가며 끝부분의 정지 화면을 연장할 수 있습니다. 진행 카드는 당시 제작한 공개 기록으로 서버 화면 녹화가 아닙니다. 설명 카드는 공개 문서의 요약·재배치입니다. 자세한 출처와 시각은 sources/provenance.json에 있습니다.

운영 서비스/실제 키/보호된 백업 자료는 새로 접근하지 않았습니다. 사례 상태는2026-10-20 공개 문서 기준이며 첫 예약 성공·전체 Mailu와 앱 복구를 완료했다고 표현하지 않습니다.

## 다시 편집하기

프로젝트의 `scripts/build_demo_video.py`를 requirements-media.txt 도구 환경에서 실행합니다. `--capture`는 공개 브라우저 촬영과 합성 시험을 새로 실행합니다. `--render`는 기존 촬영/설명 카드와 타임라인을 영상으로, `--package`는 영상 포함PPTX/대본을 만듭니다. `--verify`는 디코딩/길이/해상도/PPTX 파일·원본 해시를 확인하고 ZIP을 만듭니다. 기본값은 모든 단계입니다. 원본 WebM은 무음이며 속도 변경 없이 발췌합니다. 과거 서버 실행을 재연한 것처럼 소개하지 마세요.

출처: https://github.com/sungbeom78/dev-process-light · https://github.com/edubeom/smb-backup-lab
''')
    (OUT/'index.html').write_text('''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>smb-backup-lab 시연 영상</title><style>*{box-sizing:border-box}body{margin:0;background:#102c41;color:#edf6fa;font:18px/1.65 "Noto Sans CJK KR",sans-serif;padding:28px 5vw}main{max-width:1400px;margin:auto}h1{font-size:32px;margin:0}p{color:#b9d3df}.row{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}button,a{background:#1d4d59;border:1px solid #73d7c5;border-radius:10px;color:#f0fbf8;padding:10px 18px;font:inherit;text-decoration:none;cursor:pointer}button[aria-pressed=true]{background:#77dec9;color:#102c41}video{display:block;width:100%;aspect-ratio:16/9;background:#091a28;border-radius:14px}.links{margin-top:25px}.links a{display:inline-block;margin:6px 8px 6px 0}small{color:#b9d3df}</style>
<main><h1>dev-process-light × smb-backup-lab</h1><p>90초 시연 두 편 · 3분 합본 · 무음 / 한국어 자막 포함</p>
<div class="row"><button data-src="demo-01-records.mp4" aria-pressed="true">시연 1 · 기록 읽기 (90초)</button><button data-src="demo-02-tests.mp4" aria-pressed="false">시연 2 · 실제 시험 (90초)</button><button data-src="demo-combined.mp4" aria-pressed="false">합본 (3분)</button></div>
<video controls playsinline preload="metadata" poster="previews/demo-01-records.png" src="demo-01-records.mp4"></video>
<p id="status" aria-live="polite">시연 1 · 재생 버튼을 누르세요. 발표자는 직접 설명합니다.</p>
<small>이번 공개 화면·합성 시험 촬영과 과거 진행 카드를 장면 안에 표시했습니다. 실제 운영 백업과 키 입력 장면은 없습니다.</small>
<div class="links"><a href="smb-backup-lab-demo-embedded.pptx">영상 포함 PPTX</a><a href="smb-backup-lab-demo-kit.zip">전체 ZIP</a><a href="narration-and-edit-cues.md">대본과 편집표</a><a href="README.md">사용 안내</a></div></main>
<script>const v=document.querySelector('video');document.querySelectorAll('button[data-src]').forEach(b=>b.addEventListener('click',()=>{v.pause();v.src=b.dataset.src;v.poster='';v.load();document.querySelectorAll('button').forEach(x=>x.setAttribute('aria-pressed',String(x===b)));document.querySelector('#status').textContent=b.textContent+' · 재생 버튼을 누르세요.'}));v.addEventListener('error',()=>document.querySelector('#status').textContent='이 브라우저에서 재생되지 않으면 MP4를 내려받아 동영상 플레이어로 열어 주세요.');</script></html>''')
    print('Embedded-video PPTX and narration/edit guide saved',flush=True)


def verify():
    checks={}; video_meta={}
    for name,expected in [('demo-01-records',90),('demo-02-tests',90),('demo-combined',180)]:
        path=OUT/f'{name}.mp4'
        probe=subprocess.run([FFMPEG,'-hide_banner','-i',str(path)],capture_output=True,text=True).stderr
        duration=re.search(r'Duration: (\d+):(\d+):(\d+\.\d+)',probe)
        if not duration:raise RuntimeError('No duration in video metadata')
        seconds=int(duration[1])*3600+int(duration[2])*60+float(duration[3])
        assert abs(seconds-expected)<0.15,(name,seconds)
        assert '1920x1080' in probe and '24 fps' in probe and 'h264' in probe and 'aac' in probe
        call([FFMPEG,'-v','error','-i',path,'-f','null','-'])
        video_meta[name]={'seconds':seconds,'expected_seconds':expected,'size':[1920,1080],'fps':24,'full_decode':'passed','silent_audio':True,'burned_korean_captions':True}
        print(f'Full decode passed: {name} ({seconds}s)',flush=True)
    prs=Presentation(OUT/'smb-backup-lab-demo-embedded.pptx');assert len(prs.slides)==26
    with zipfile.ZipFile(OUT/'smb-backup-lab-demo-embedded.pptx') as z:
        assert z.testzip() is None
        media=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(media)==2
        expected={sha(OUT/'demo-01-records.mp4'),sha(OUT/'demo-02-tests.mp4')}
        assert {hashlib.sha256(z.read(n)).hexdigest() for n in media}==expected
        assert len([n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)])==26
        for n in [22,23]:
            assert 'videoFile' in z.read(f'ppt/slides/slide{n}.xml').decode()
        for name in z.namelist():
            if name.endswith('.xml'):check_text(z.read(name).decode())
    # Frame samples cover all fourteen scenes; these are the exact burned output.
    timeline=json.loads((OUT/'timeline.json').read_text())
    thumbs=[]
    for s in timeline['scenes']:
        frame=OUT/'previews'/f'scene-{s["number"]:02}.jpg'
        call([FFMPEG,'-y','-loglevel','error','-ss',str(s['start']+s['duration']/2),'-i',OUT/'demo-combined.mp4','-frames:v','1',frame])
        thumbs.append(frame)
    sheet=Image.new('RGB',(1440,4*300),'#0d2233');draw=ImageDraw.Draw(sheet)
    for i,path in enumerate(thumbs):
        im=Image.open(path);im.thumbnail((470,264));x=(i%3)*480;y=(i//3)*300
        # Five rows for fourteen shots.
        if y+300>sheet.height:
            larger=Image.new('RGB',(1440,y+300),'#0d2233');larger.paste(sheet,(0,0));sheet=larger;draw=ImageDraw.Draw(sheet)
        sheet.paste(im,(x,y));draw.text((x+8,y+267),f'{i+1:02}  {timeline["scenes"][i]["title"]}',font=ImageFont.truetype(FONT,18,index=1),fill='white')
    sheet.save(OUT/'previews/contact-sheet.jpg')
    checks.update(video_decode_duration_resolution='passed',pptx_slides=26,pptx_embedded_video_count=2,pptx_video_hashes='passed',pptx_notes=26,
                  powerpoint_application_playback='not tested; app unavailable',privacy_text_scan='passed',service_operations_performed=False,external_upload_performed=False)
    for path in OUT.rglob('*'):
        if path.is_file() and path.suffix in ['.txt','.md','.html','.json','.srt','.ass']:
            check_text(path.read_text())
    check_player()
    checks['source_provenance']=json.loads((OUT/'sources/provenance.json').read_text())
    files={str(p.relative_to(OUT)):sha(p) for p in sorted(OUT.rglob('*')) if p.is_file() and p.name not in ['manifest.json','smb-backup-lab-demo-kit.zip']}
    put_json(OUT/'manifest.json',{'created_at':now(),'videos':video_meta,'checks':checks,'sha256':files})
    target=OUT/'smb-backup-lab-demo-kit.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=5) as z:
        for path in sorted(OUT.rglob('*')):
            if path.is_file() and path != target:z.write(path,Path('smb-backup-lab-demo')/path.relative_to(OUT))
    with zipfile.ZipFile(target) as z:assert z.testzip() is None
    print('Verified and packaged: '+str(OUT),flush=True)


def check_player():
    results=[];errors=[]
    names=['demo-01-records','demo-02-tests','demo-combined']
    with sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1600,'height':1100})
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.goto((OUT/'index.html').as_uri())
        for i,wanted in enumerate([90,90,180]):
            page.locator('button[data-src]').nth(i).click()
            page.wait_for_function('Number.isFinite(document.querySelector("video").duration)',timeout=20000)
            meta=page.locator('video').evaluate('v=>({duration:v.duration,width:v.videoWidth,height:v.videoHeight})')
            assert abs(meta['duration']-wanted)<.2 and (meta['width'],meta['height'])==(1920,1080)
            page.locator('video').evaluate('v=>v.play()');page.wait_for_timeout(800)
            assert page.locator('video').evaluate('v=>v.currentTime')>0
            page.locator('video').evaluate('v=>{v.pause();v.currentTime=10}');page.wait_for_timeout(400)
            assert abs(page.locator('video').evaluate('v=>v.currentTime')-10)<.1
            results.append({'item':i+1,**meta,'play_and_seek':'passed'})
        assert not errors,errors
        page.screenshot(path=str(OUT/'previews/player-browser.png'))
        browser.close()
    put_json(OUT/'browser-check.json',{'results':results,'page_errors':errors,'video_sha256':{name:sha(OUT/f'{name}.mp4') for name in names}})
    print('Browser play and seek passed for all three MP4s',flush=True)


def finalize_review():
    """Seal additional human/agent frame review and browser evidence after verify.

    Existing verified files must be byte-identical; this does not replace a media
    verification pass when an encoded video or its source changes.
    """
    manifest=json.loads((OUT/'manifest.json').read_text())
    for name,digest in manifest['sha256'].items():
        if name in ['browser-check.json','previews/player-browser.png']:continue
        if sha(OUT/name)!=digest:raise RuntimeError('Verified file changed before final review: '+name)
    browser=json.loads((OUT/'browser-check.json').read_text())
    assert not browser['page_errors'] and len(browser['results'])==3
    assert all(r['play_and_seek']=='passed' for r in browser['results'])
    assert browser['video_sha256']=={name:sha(OUT/f'{name}.mp4') for name in ['demo-01-records','demo-02-tests','demo-combined']}
    manifest['checks']['browser_player']=browser
    manifest['checks']['visual_review']='14-scene contact sheet and full-size samples inspected; captions/source labels fit'
    manifest['checks']['reviewed_at']=now()
    manifest['sha256']={str(p.relative_to(OUT)):sha(p) for p in sorted(OUT.rglob('*')) if p.is_file() and p.name not in ['manifest.json','smb-backup-lab-demo-kit.zip']}
    put_json(OUT/'manifest.json',manifest)
    target=OUT/'smb-backup-lab-demo-kit.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=5) as z:
        for path in sorted(OUT.rglob('*')):
            if path.is_file() and path!=target:z.write(path,Path('smb-backup-lab-demo')/path.relative_to(OUT))
    with zipfile.ZipFile(target) as z:assert z.testzip() is None
    print('Final review evidence included in ZIP',flush=True)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--recapture-public',action='store_true',help='Repeat only public GitHub and lecture recordings')
    p.add_argument('--recapture-tests',action='store_true',help='Repeat only synthetic test capture and explanation cards')
    p.add_argument('--check-player',action='store_true',help='Verify browser play and seeking without re-encoding')
    p.add_argument('--refresh-identity',action='store_true',help='Refresh only identity art and provenance; keep test recordings')
    p.add_argument('--finalize-review',action='store_true',help='After inspecting samples, include unchanged-media browser/review evidence')
    for name in ['capture','render','package','verify']:p.add_argument('--'+name,action='store_true')
    args=p.parse_args();selected=[name for name in ['capture','render','package','verify'] if getattr(args,name)]
    if args.recapture_public:
        capture(public_only=True)
        return
    if args.recapture_tests:
        capture(tests_only=True)
        return
    if args.finalize_review:
        finalize_review()
        return
    if args.check_player:
        check_player()
        return
    if args.refresh_identity:
        refresh_identity()
        return
    for name in selected or ['capture','render','package','verify']:globals()[name]()


if __name__=='__main__':main()
